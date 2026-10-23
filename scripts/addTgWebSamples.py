"""
TG-WEB「サンプル1（重み=1）」として25問を採用シートへ追加・加算する。

方針:
- 解答テキストが既存グループの「採用解答」と前方一致 or 完全一致 → そのグループに +1票
- 上記でマッチしなければ → 新規グループとして末尾に追加
- 票が増えたグループは confidence/rank/総票数/加重票数/得票内訳/出典ファイル数を再計算
- 詳細シートは触らない（採用シートのみ）
"""
from __future__ import annotations
import math
import re
from pathlib import Path
from openpyxl import load_workbook

XLSX = Path(__file__).resolve().parents[1] / "TG-WEB（正解確定版）のコピー.xlsx"

# (カテゴリ, シート名)
CATEGORIES = [
    ("gengo", "言語_採用"),
    ("higengo", "非言語_採用"),
    ("zukei", "図形_採用"),
]

# サンプル1（重み=1）として投入するデータ
QUERIES: dict[str, list[tuple[str, str, str]]] = {
    "gengo": [
        ("観光旅行に出ることは",
         "旅は、未知の領域に足を踏み入れることで何らかのものと向き合う精神的意義が大きい。",
         "観光旅行に出ることは、旅に出ることとは違う。観光旅行はガイドブックに紹介された場所や、大勢の人が何度も見聞した有名なところを訪ねることである。それに対して旅に出るというのは、未知の場所に足を踏みいれることであり、知っている範囲を超えて、勇気を持って新しい場所に向かうことである。それは、肉体的、または空間的な意味合いだけでなく、精神的な部分も含まれるのである。むしろ、精神的部分が強いと言ってもよい。…生まれたばかりの子どもにとっても、世界は異質なものであふれかえっている。…そういう意味では、子どももまた究極の旅人だといえるかもしれない。"),
        ("財務諸表とは",
         "各企業が独自の方法や解釈で決算書を作成",
         "財務諸表とは、複式簿記に基づく貸借対照表ならびに損益計算書、さらに株主資本等変動計算書などによって構成され、企業が株主、出資者、あるいは利害関係者などに対して、一定期間の経営成績、財務状態を明示するために作成される報告書で、一般的に決算書ともいわれる。…決算書とは、要するに企業の「通信簿」なのである。社会一般に報告するのが目的なので、その形式は決められている。（　　　）したら、比較のしようがなくなってしまう。内部資料ではないので、社会に通用する内容であることが求められる。"),
        ("江戸や明治の時代には",
         "A・B・D・E・H",
         "A〜Iには「地域社会」もしくは「国や県（国家）」が入る。「地域社会」が入るもののみをあげているのはどれか。\n江戸や明治の時代には、それぞれの（A）のコミュニティが、農業や漁業に携わりながら自分たちで身近な水を管理していました。…川を利用するのも、川の災害から自身を守るのも、（B）の役割でした。\n1896年、河川法ができて、（C）による川の管理が始まりました。…（D）と川との関わりが薄れ始めました。…しかし、（E）の側からすると、川は（F）が管理するものとして、自分たちとの関わりがなくなっていったのです。…洪水の時には（G）が手当てをしてくれるようになりました。（H）がそれを歓迎したということもあります。…（I）がやってくれたほうがありがたいと、地元でも判断されるようになったからです。"),
        ("人間は互いに干渉し合って",
         "互いの関係性を保つためのルール",
         "人間は互いに干渉し合って生きている存在である。そして、人間同士が共存していくためには、（　　）が必要となるのである。中世においての人間関係は、謂わば忠誠という法に基づいて成り立っていたと言われる。…対して現代社会においての人間関係は、契約に基づいて成り立っている。…契約というものは、信頼とは正反対の意味を持つ。…ようするに、人は裏切るという前提のもと成り立っている関係だからである。"),
        ("鉱石ラジオ",
         "人生を輝かせてくれる喜びはある",
         "形の表現の面では大変魅力的な鉱石ラジオなのだが、イヤホンから聞こえる音は小さいことが多く、よほど耳を澄ませなければ聞けないときもある。それを性能の悪さと取る方もいるだろうが、ぼくはそうは思わない。…ぼくたちが多忙で自分を見失いがちな日常の中にも、（　　　）。鉱石ラジオはそんなことを教えてくれる。"),
        ("ドルトン",
         "C→E→D→A→B",
         "A〜Eの文を意味が通るように並べ替えたものはどれか。\nA：その後、アトムは、分割できない粒子ではなく、構造を持つことが、原子からの光放射や電子の発見によって明らかになってきた。\nB：真空状態での原子と光の反応実験が行われるようになったからである。\nC：1803年、化学者のドルトンは、化学実験によって発見された定比例の法則や倍数比例の法則を理解するために、「原子論」を提唱した。\nD：彼は、ギリシャのデモクリトスにちなんで、原子を「アトム」と名付けた。\nE：各元素は、おのおの決まった「原子」から成り、「化合物」は原子が簡単な比で結合した「分子」からできていると考えれば、簡単に説明できるというものである。"),
        ("模倣から創造",
         "D→C→A→E→B",
         "A：これらは、創造性の能力が高まり、新しい発見をする準備が整いつつあることの証しでもある。\nB：そこで、自分の力で何か新しい最先端の研究がしたくなり、実際に研究を始めてみると、難しさと同時に手ごたえも得られるようになる。\nC：それは、最新の科学雑誌に目を通すと、何が新しく分かったか、何がまだ分かっていないかが頭に入ってくるようになる時期であるともいえる。\nD：模倣から創造へ移行する段階は、普通の科学教育の場合、大学卒業から大学院にかけて訪れる。\nE：つまり、その人が真理の大海に漕ぎだす時期でもある。"),
        ("展覧会とは不思議",
         "E→C→A→B→D",
         "A：つまり、個々の芸術作品は、特定のテーマを掲げた展覧会の中に組み込まれ、他の作品と相互に比較されることで、はじめて表現としての完結を見る性質のものではないのだ。\nB：多くの場合、展示は傾向別、年代順などの作品の外にある基準に依拠している。\nC：アトリエにおける芸術家の孤独な作業として生まれた作品は、それぞれが独立した価値を持ち、単独で鑑賞されるべきものであり、必ずしも多くの作品を同じ場所に集めて、同時に鑑賞させる形式を必要としないはずである。\nD：悪くいえば、そうした基準は事後的・便宜的につくられたものであって、自立した作品の鑑賞の妨げになりかねない。\nE：展覧会とは不思議なものである。我々は展覧会に慣れているのでさほど疑問には思わないが、美術作品の本来のあり方に照らせば、不思議な鑑賞形式といわざるをえない。"),
        ("やまびこあいさつ",
         "よい印象を与えるはずのサービスが、かえって客を不快にさせることもある。",
         "コンビニエンスストア、大型新古書店をはじめ、多くの小売店で採用されている「やまびこあいさつ」というものがある。…大学生にアンケートをとってみると「やまびこあいさつ」について「心がこもっていない」「反射的にさけんでいるだけ」「パーソナルな応対ではなく、適当にこなしている感じ」と、手厳しい結果が出た。…「ありがとうございました」のこだまを浴びせかけるのは、むしろ逆効果ということになる。"),
        ("りんごの味",
         "古典の享受にあたって、どの作品を読めばよいかは決まっていない。",
         "りんごの味がどんなものであるかを知るには、りんごを食べることが先決で、食べもせずにその味について語ることは観念の遊びである。古典についても同様で、まずそれをみずから読むという実行こそが基本である。…何を選んで読むかということもよく問題にされるが、これにはあまりこだわる必要はない。自分の出会った作品、興味を覚えたものから読み始めればいいわけで、これを読むとよいと外から決めつけられるのは好ましくない。"),
        ("合計特殊出生率",
         "幼保一体化は待機児童解消が大きな目的である。",
         "女性1人が生涯に産む子どもの推定人数を示す合計特殊出生率は、ここ数年1.35前後である。…少子化に歯止めをかける対策は明白だ。女性が安心して子どもを産み育てられる環境をどう整えるかである。この問題解決の一案として、幼稚園と保育所の一元化が挙げられる。保育所に入れない待機児童の解消を目指し、その大多数を占める0—2歳児の保育の場を増やすことが主なねらいである。"),
        ("統計を使う学問",
         "学問の世界全体は、その全体性において解体寸前の危機的状況、アイデンティティの喪失状態にある。",
         "方法論的には、統計を使う学問は大体局所より全体を見ようとしています。…いま学問の世界全体は、細分化、細密化がどんどん進んでいくことによって、全体としては、統合が失われ、解体しつつあるといってもいいのです。…全体性とかインテグレーション（integration 統合）という角度から見ようとすると、解体寸前の危機的状況にあるわけです。いま学問は、…アイデンティティ喪失状態にあるといってもいいわけです。"),
        ("日本座敷を一つの墨絵",
         "空間を遮蔽して生ずる陰翳の世界は、無気味な静かさを持つ東洋の神秘だといえる。",
         "もし日本座敷を一つの墨絵にたとえるなら、障子は墨色の最も淡い部分であり、床の間は最も濃い部分である。…思うに西洋人のいう「東洋の神秘」とは、かくの如き暗がりが持つ無気味な静かさを指すのであろう。…われらの祖先の天才は、虚無の空間を任意に遮蔽しておのづから生ずる陰翳の世界に、いかなる壁画や装飾にも優る幽玄味を持たせたのである。"),
        ("生まれ落ちた時",
         "成熟とは、他者の都合を考え、尊重することである。",
         "人がこの世に生まれ落ちた時、そこに立ち現れるのは自己と他者が未分化の世界である。…他者との関わりの中で、自分の欲望と他者の欲望が必ずしも調和しないこと、他者には他者の都合があり、時にはそれを尊重すべきであると学習することを、人は一人の人間としての価値を所有する大人になる、つまり成熟すると言う。"),
    ],
    "higengo": [
        ("正八面体の展開図",
         "アとエ",
         "図のような展開図を組み立てて、正八面体を作ったとき、キ以外で面αと接している面はどれか。（展開図：ア-イ-ウ-エ-オ-カ-キ-α が帯状に並ぶ）"),
        ("安藤君のアパート",
         "3号室",
         "住所を頼りに安藤君のアパートを訪ねたが、このアパートには1号室から6号室まで6部屋あり、表札がないため、どれが安藤君の部屋か分からなかった。そこで、通りすがりの近所の人に聞くと、「井上さんと上田さんは部屋番号が2つ違い、遠藤さんと岡田さんの部屋番号を足すと7になります。5号室は誰も住んでいませんよ」と答えた。安藤君は何号室に住んでいるか。"),
        ("果物の好み",
         "みかんとなしが両方好きな人がいる。",
         "ある高校の生徒に果物の好みについてアンケートをとった。\nア：なしが好きな人の中にりんごを好きな人がいる。\nイ：なしが好きな人は必ずももが好きである。\nウ：りんごが好きな人は必ずみかんが好きである。\nこのとき、確実にいえることは次のうちどれか。"),
        ("AはCより3歳年上",
         "Cは2番目に年上である。",
         "A〜Eの年齢について次のア〜オのことが分かっている。\nア：AはCより3歳年上である。\nイ：BとAは6歳違いである。\nウ：CとDは5歳違いである。\nエ：DはBより2歳年下である。\nオ：EはBと2歳違いであり、Cと1歳違いである。\nこれらのことから確実にいえるのはどれか。"),
        ("Pは1位だ",
         "8通り",
         "P、Q、R、Sの4人が水泳の競争をすることになり、その結果を4人が次のように予想した。\n予想①：「Pは1位だ。」\n予想②：「Qは2位だ。」\n予想③：「Rは3位ではない。」\n予想④：「Sは3位ではない。」\n競争の結果、同じ順位の者はなく、予想①と②がはずれ、予想③と④が当たった。このとき、4人の順位として考えられるのは何通りか。"),
        ("赤、白、青、黒のTシャツ",
         "Dは青のTシャツを着ていた。",
         "A、B、C、Dは赤、白、青、黒のTシャツのいずれかを着ており、同じ色を着ていた者はいない。各人は他の3人のうち2人と出会った。\nア：Aは赤のTシャツを着た者と出会った。\nイ：BはCと出会ったが、青のTシャツを着た者と出会わなかった。\nウ：Cは白のTシャツを着た者と出会わなかった。\nエ：赤のTシャツを着た者は青のTシャツを着た者とは出会わなかった。"),
        ("Aの家から北に400m",
         "AD間の距離はBC間の距離の2倍である。",
         "Aの家から北に400m行き、東に400m行くとBの家がある。Cの家はA、B、Dの家から等距離にある。Dの家から南に400m行き、東に400m行くとBの家がある。次のうち正しいものはどれか。"),
    ],
    "zukei": [
        ("正方形の紙を2回三つ折り",
         "E",
         "正方形の紙を2回三つ折りにする。その後、黒色部分（折りたたみ後の右下角の小さな正方形）を切り取ったとき、切り取った後の紙をもとのように広げた図形として正しいものはどれか。"),
        ("立方体の展開図に余分な正方形",
         "3通り",
         "次の図は立方体の展開図に余分な正方形を2個追加した図形である。追加した2個の正方形として考えられる組合せは何通りあるか。（8マスの図形、配置は中段4マス＋上段左2マス＋下段の2マス）"),
        ("T字型",
         "24枚",
         "次の図のような図形（T字型／凸型：上部10×5、下部20×5）をすき間なく敷き詰めて正方形を作るには、最低何枚必要か。"),
        ("黒丸2個・白丸2個",
         "A",
         "次の展開図（黒丸2個・白丸2個付き）を組み立てた立方体として正しいものは、A〜Eのうちどれか。"),
    ],
}

# 採用シートのカラムインデックス（1-based）
COL = {
    "no": 1, "q": 2, "a": 3, "conf": 4, "rank": 5, "verify": 6,
    "total": 7, "weighted": 8, "breakdown": 9, "files": 10, "key": 11, "group": 12,
}

PUNCT_RE = re.compile(r"[\s　,，、。\.\(\)（）\[\]【】「」『』:：;；!！?？/／\-—–→\->・]")


def norm_q(s) -> str:
    if s is None:
        return ""
    return PUNCT_RE.sub("", str(s))


def norm_a(s) -> str:
    if s is None:
        return ""
    return PUNCT_RE.sub("", str(s)).lower()


def calc_rank(weighted_for_adopted: float, total_weighted: float) -> tuple[str, int]:
    if total_weighted <= 0:
        return "N", 0
    agreement = weighted_for_adopted / total_weighted
    depth = math.sqrt(min(1.0, total_weighted / 3.0))
    conf = round(agreement * depth * 100)
    if conf >= 90:
        rank = "A"
    elif conf >= 70:
        rank = "B"
    elif conf >= 50:
        rank = "C"
    else:
        rank = "D"
    return rank, conf


def parse_breakdown(s: str) -> list[tuple[str, float]]:
    """得票内訳「A:3.0, B:2.0」をパースして [(text, weight), ...] を返す。"""
    if not s:
        return []
    parts = []
    for chunk in str(s).split(","):
        chunk = chunk.strip()
        if not chunk:
            continue
        # 末尾の「:数値」を切り出す
        m = re.match(r"^(.*):([\d.]+)$", chunk)
        if not m:
            continue
        text = m.group(1).strip()
        try:
            w = float(m.group(2))
        except ValueError:
            continue
        parts.append((text, w))
    return parts


def serialize_breakdown(items: list[tuple[str, float]]) -> str:
    items_sorted = sorted(items, key=lambda x: -x[1])
    return ", ".join(f"{t}:{w:g}" for t, w in items_sorted)


def find_match(rows: list[dict], q_kw: str, ans: str) -> dict | None:
    """問題文HIT × 解答テキスト一致の最初を返す。なければ None。"""
    nq = norm_q(q_kw)
    na = norm_a(ans)
    for r in rows:
        if nq not in norm_q(r["q"]):
            continue
        ra = norm_a(r["a"])
        if not ra or not na:
            continue
        # 解答テキストが完全一致 / 一方が他方の前方部分（5文字以上のとき）
        if ra == na:
            return r
        if len(na) >= 5 and (ra.startswith(na) or na.startswith(ra)):
            return r
        # 共通先頭が解答長の8割以上
        common = 0
        for x, y in zip(ra, na):
            if x == y:
                common += 1
            else:
                break
        if common >= max(len(ra), len(na)) * 0.8 and common >= 5:
            return r
    return None


def main():
    wb = load_workbook(XLSX)
    summary = []

    for cat, sheet_name in CATEGORIES:
        ws = wb[sheet_name]
        if cat not in QUERIES:
            continue

        # 既存行を読み込み
        rows = []
        for r_idx in range(2, ws.max_row + 1):
            no = ws.cell(r_idx, COL["no"]).value
            if no is None or str(no).strip() == "":
                continue
            rows.append({
                "row": r_idx,
                "no": int(no),
                "q": ws.cell(r_idx, COL["q"]).value or "",
                "a": ws.cell(r_idx, COL["a"]).value or "",
                "conf": ws.cell(r_idx, COL["conf"]).value,
                "rank": ws.cell(r_idx, COL["rank"]).value,
                "verify": ws.cell(r_idx, COL["verify"]).value,
                "total": ws.cell(r_idx, COL["total"]).value,
                "weighted": ws.cell(r_idx, COL["weighted"]).value,
                "breakdown": ws.cell(r_idx, COL["breakdown"]).value,
                "files": ws.cell(r_idx, COL["files"]).value,
                "key": ws.cell(r_idx, COL["key"]).value,
                "group": ws.cell(r_idx, COL["group"]).value,
            })

        max_no = max(r["no"] for r in rows)
        max_group = max(int(r["group"]) for r in rows if r["group"] not in (None, ""))

        added = 0
        boosted = 0

        for q_kw, expected_a, full_q in QUERIES[cat]:
            target = find_match(rows, q_kw, expected_a)
            if target:
                # 既存グループに +1票（重み=1）
                cur_total = int(target["total"]) if target["total"] not in (None, "") else 0
                cur_weighted = float(target["weighted"]) if target["weighted"] not in (None, "") else 0.0
                cur_files = int(target["files"]) if target["files"] not in (None, "") else 0
                breakdown = parse_breakdown(target["breakdown"] or "")
                # 採用解答に対応する得票内訳を +1
                # （内訳テキストが採用解答と表記ゆれしていてもよい：完全一致 or 前方一致 or 一方が短縮形）
                ans_norm_target = norm_a(target["a"])

                def is_same_answer(txt: str) -> bool:
                    nx = norm_a(txt)
                    if not nx or not ans_norm_target:
                        return False
                    if nx == ans_norm_target:
                        return True
                    short, long = sorted([nx, ans_norm_target], key=len)
                    if len(short) >= 5 and long.startswith(short):
                        return True
                    return False

                matched_idx = None
                for i, (txt, w) in enumerate(breakdown):
                    if is_same_answer(txt):
                        matched_idx = i
                        break
                if matched_idx is not None:
                    txt, w = breakdown[matched_idx]
                    breakdown[matched_idx] = (txt, w + 1.0)
                else:
                    breakdown.append((str(target["a"]), 1.0))

                new_total = cur_total + 1
                new_weighted = cur_weighted + 1.0
                # 採用解答に紐づく加重票を再計算（表記ゆれ許容）
                adopted_weighted = sum(w for txt, w in breakdown if is_same_answer(txt))
                rank, conf = calc_rank(adopted_weighted, new_weighted)

                ws.cell(target["row"], COL["total"]).value = new_total
                ws.cell(target["row"], COL["weighted"]).value = new_weighted
                ws.cell(target["row"], COL["breakdown"]).value = serialize_breakdown(breakdown)
                ws.cell(target["row"], COL["files"]).value = cur_files + 1
                ws.cell(target["row"], COL["conf"]).value = conf
                ws.cell(target["row"], COL["rank"]).value = rank
                # 要検証は信頼度が D 未満時のみ立てる方針（ここでは触らない）
                boosted += 1
                summary.append((cat, "BOOST", target["no"], rank, conf, target["a"]))
            else:
                # 新規追加
                max_no += 1
                max_group += 1
                rank, conf = calc_rank(1.0, 1.0)  # 1票だけ → conf=58, rank=C
                new_row_idx = ws.max_row + 1
                ws.cell(new_row_idx, COL["no"]).value = max_no
                ws.cell(new_row_idx, COL["q"]).value = full_q
                ws.cell(new_row_idx, COL["a"]).value = expected_a
                ws.cell(new_row_idx, COL["conf"]).value = conf
                ws.cell(new_row_idx, COL["rank"]).value = rank
                ws.cell(new_row_idx, COL["verify"]).value = ""
                ws.cell(new_row_idx, COL["total"]).value = 1
                ws.cell(new_row_idx, COL["weighted"]).value = 1.0
                ws.cell(new_row_idx, COL["breakdown"]).value = f"{expected_a}:1"
                ws.cell(new_row_idx, COL["files"]).value = 1
                ws.cell(new_row_idx, COL["key"]).value = norm_a(expected_a)
                ws.cell(new_row_idx, COL["group"]).value = max_group
                added += 1
                summary.append((cat, "NEW", max_no, rank, conf, expected_a))

        print(f"[{cat:8s}] +{boosted} boosted / +{added} new")

    wb.save(XLSX)
    print("\n===== 詳細 =====")
    for cat, kind, no, rank, conf, ans in summary:
        ans_disp = ans if len(str(ans)) <= 50 else str(ans)[:50] + "…"
        print(f"  [{cat:8s}] {kind:5s}  No.{no:5d}  rank={rank} conf={conf:3d}  {ans_disp}")


if __name__ == "__main__":
    main()
